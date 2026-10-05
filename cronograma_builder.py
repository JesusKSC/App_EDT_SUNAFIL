"""
cronograma_builder.py — Sección 3 (Cronograma) de la app "Checklist de TDR — OCR (Gemini) -> Excel".

Convierte la EDT (28 paquetes de edt_builder.py) en un cronograma vertical "tipo MS Project":
  Proyecto (1) > Entregable (1.x) > Paquete (1.x.y) > Actividad con verbo (1.x.y.z) / Hito

Genera CRONOGRAMA_SUNAFIL.xlsx con hojas: Resumen, Cronograma, Gantt implementación (diario),
Gantt general (mensual) y Costos. Las fechas y costos del Excel son FÓRMULAS: si cambias la fecha
de firma (celda amarilla) o una duración, todo se recalcula. No usa Gemini.

Uso:
    from cronograma_builder import construir_cronograma, excel_cronograma
    cr = construir_cronograma(fecha_firma, plazo_activacion=14, monto=1614806.00, plazo_servicio=731)
    xlsx_bytes = excel_cronograma(cr)
    cr["tabla"] (DataFrame para st.dataframe), cr["metricas"], cr["cumple_activacion"]
"""
import datetime as dt
import io

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

from edt_builder import ENTREGABLES, PAQUETES, PROYECTO, SUBTITULO

# ----------------------------------------------------------------------------------------
# 1. ACTIVIDADES POR PAQUETE (verbos). dur = días calendario; "SERV" = plazo del servicio;
#    "SERV+10" = plazo del servicio + 10 días (entregables y pagos del último mes).
#    preds = claves de las predecesoras. base = de dónde sale el plazo.
#    costo = estimación referencial (S/); None = recibe el saldo del monto contractual.
# ----------------------------------------------------------------------------------------
EQ, TDR = "Estimación del equipo", "Plazo del TDR o del contrato"
ACT = [
    # paquete, clave, actividad, dur, preds, costo, responsable, base
    ("1.1.1", "LIN1", "Reservar la numeración y crear las líneas en la plataforma", 5, ["FIRMA"], 3000, "Líder de Operaciones", EQ),
    ("1.1.1", "LIN2", "Activar las líneas y hacer pruebas de tráfico", 4, ["POR3", "LIN1"], 4000, "Líder de Operaciones", EQ),
    ("1.1.1", "OPER", "Operar el servicio móvil (voz, datos y SMS)", "SERV", ["INICIO"], None, "Líder de Operaciones", TDR),
    ("1.1.2", "VOZ", "Configurar el plan de voz y los perfiles de llamadas", 2, ["LIN1"], 1000, "Líder de Operaciones", EQ),
    ("1.1.3", "DAT", "Configurar las bolsas de datos por gama con las mejoras", 2, ["LIN1"], 1000, "Líder de Operaciones", EQ),
    ("1.1.4", "SMS", "Habilitar SMS y servicios suplementarios", 2, ["LIN1"], 800, "Líder de Operaciones", EQ),
    ("1.1.5", "ROA", "Habilitar el roaming en las líneas autorizadas", 2, ["LIN1"], 800, "Líder de Operaciones", EQ),
    ("1.2.1", "GA1", "Adquirir los equipos Gama A", 7, ["FIRMA"], 24000, "Líder de Logística", EQ),
    ("1.2.1", "GA2", "Entregar los equipos Gama A con acta", 4, ["ENR"], 400, "Líder de Logística", EQ),
    ("1.2.2", "GB1", "Adquirir los equipos Gama B", 7, ["FIRMA"], 60000, "Líder de Logística", EQ),
    ("1.2.2", "GB2", "Entregar los equipos Gama B con acta", 4, ["ENR"], 1500, "Líder de Logística", EQ),
    ("1.2.3", "GC1", "Adquirir los equipos Gama C", 7, ["FIRMA"], 442800, "Líder de Logística", EQ),
    ("1.2.3", "GC2", "Distribuir los equipos Gama C a la sede y las dependencias", 4, ["ENR"], 11000, "Líder de Logística", EQ),
    ("1.2.4", "ACC", "Entregar los equipos de respaldo y los accesorios", 4, ["ENR"], 12000, "Líder de Logística", EQ),
    ("1.3.1", "MDM1", "Aprovisionar la consola MDM en la nube", 3, ["FIRMA"], 25000, "Especialista MDM", EQ),
    ("1.3.1", "MDM2", "Configurar las políticas de seguridad y de aplicaciones", 3, ["MDM1"], 3000, "Especialista MDM", EQ),
    ("1.3.2", "WEB1", "Habilitar el portal web 24x7 y crear los usuarios", 3, ["FIRMA"], 2000, "Especialista MDM", EQ),
    ("1.3.2", "WEB2", "Validar los reportes de consumo con la OTIC", 2, ["WEB1", "LIN1"], 800, "Especialista MDM", EQ),
    ("1.4.1", "COB1", "Verificar la cobertura en la sede central y las dependencias", 5, ["FIRMA"], 6000, "Ingeniero de Radio", EQ),
    ("1.4.2", "IND1", "Medir la señal indoor e identificar los puntos críticos", 5, ["INICIO"], 4000, "Ingeniero de Radio", EQ),
    ("1.4.2", "IND2", "Implementar la solución indoor en las sedes que lo requieran", 30, ["IND1"], 60000, "Ingeniero de Radio", TDR),
    ("1.4.2", "IND3", "Emitir el informe de mediciones posterior a la solución", 2, ["IND2"], 1000, "Ingeniero de Radio", EQ),
    ("1.5.1", "FIRMA", "Hito: Contrato suscrito", 0, [], 0, "SUNAFIL / Contratista", TDR),
    ("1.5.1", "KICK", "Realizar la reunión de inicio con la OTIC", 1, ["FIRMA"], 500, "Gestor de Servicio (PM)", EQ),
    ("1.5.1", "PLAN", "Elaborar y aprobar el plan de trabajo", 2, ["KICK"], 1500, "Gestor de Servicio (PM)", EQ),
    ("1.5.2", "INF", "Elaborar y presentar los entregables mensuales", "SERV+10", ["INICIO"], 24000, "Gestor de Servicio (PM)", TDR),
    ("1.5.3", "DOC", "Reunir y presentar la documentación para suscribir el contrato", 5, ["-FIRMA"], 2000, "Gestor de Servicio (PM)", EQ),
    ("1.5.4", "SEG", "Realizar el seguimiento mensual y el registro de riesgos", "SERV", ["INICIO"], 13200, "Gestor de Servicio (PM)", EQ),
    ("1.6.1", "POR1", "Recibir de SUNAFIL los documentos de portabilidad", 5, ["FIRMA"], 0, "SUNAFIL (OTIC)", TDR),
    ("1.6.1", "POR2", "Presentar las solicitudes de portabilidad", 2, ["POR1"], 2000, "Líder de Operaciones", EQ),
    ("1.6.1", "POR3", "Ejecutar la portabilidad y validar la numeración", 3, ["POR2"], 3000, "Líder de Operaciones", EQ),
    ("1.6.2", "ENR", "Configurar los equipos y enrolarlos en la consola MDM", 3, ["GA1", "GB1", "GC1", "MDM2"], 5000, "Especialista MDM", EQ),
    ("1.6.3", "ACTA", "Hito: Acta de culminación de implementación (entregable 11.1)", 0,
     ["LIN2", "VOZ", "DAT", "SMS", "ROA", "GA2", "GB2", "GC2", "ACC", "WEB2", "COB1", "PLAN"], 0, "Gestor de Servicio (PM)", TDR),
    ("1.6.3", "INICIO", "Hito: Inicio del servicio", 0, ["ACTA"], 0, "SUNAFIL / Contratista", TDR),
    ("1.7.1", "MAN", "Elaborar los manuales de la consola MDM y del portal web", 5, ["INICIO"], 3000, "Especialista MDM", EQ),
    ("1.7.2", "CAP1", "Coordinar la fecha y la modalidad de la capacitación", 2, ["MAN"], 300, "Especialista MDM", EQ),
    ("1.7.2", "CAP2", "Ejecutar la capacitación y emitir las constancias", 3, ["CAP1"], 4000, "Especialista MDM", EQ),
    ("1.8.1", "MESA", "Brindar atención 24x7x365 a los usuarios", "SERV", ["INICIO"], 48000, "Soporte Técnico", TDR),
    ("1.8.2", "GAR", "Gestionar las garantías de equipos, baterías y accesorios", "SERV", ["INICIO"], 20000, "Soporte Técnico", TDR),
    ("1.8.3", "REP", "Atender averías y reponer equipos", "SERV", ["INICIO"], 30000, "Soporte Técnico", TDR),
    ("1.9.1", "CONF", "Tramitar las conformidades mensuales y los pagos", "SERV+10", ["INICIO"], 2400, "Gestor de Servicio (PM)", TDR),
    ("1.9.2", "DEV", "Devolver o transferir los equipos al término del contrato", 30, ["OPER"], 3000, "Líder de Logística", TDR),
    ("1.9.3", "FIN1", "Elaborar el informe final y obtener la conformidad final", 7, ["OPER"], 1500, "Gestor de Servicio (PM)", EQ),
    ("1.9.3", "LEC", "Registrar las lecciones aprendidas", 3, ["FIN1"], 500, "Gestor de Servicio (PM)", EQ),
    ("1.9.3", "CIERRE", "Hito: Cierre del contrato", 0, ["LEC", "CONF", "INF", "DEV"], 0, "SUNAFIL / Contratista", TDR),
]
HITO_INICIO = {"INICIO"}            # sus sucesores empiezan el mismo día
FASE_PREVIA = {"DOC"}
FASE_CIERRE = {"DEV", "FIN1", "LEC", "CIERRE"}


def _dur(d, serv):
    return serv if d == "SERV" else serv + 10 if d == "SERV+10" else d


# ----------------------------------------------------------------------------------------
# 2. CÁLCULO EN PYTHON (para la pantalla); el Excel repite el cálculo con fórmulas
# ----------------------------------------------------------------------------------------
def construir_cronograma(fecha_firma, plazo_activacion=14, monto=1614806.00, plazo_servicio=731, lineas=1145, res_edt=None):
    if isinstance(fecha_firma, dt.datetime):
        fecha_firma = fecha_firma.date()
    por_clave = {a[1]: a for a in ACT}
    ini, fin = {}, {}

    def calc(k):
        if k in fin:
            return
        paq, _, nom, d, preds, *_ = por_clave[k]
        d = _dur(d, plazo_servicio)
        if k == "FIRMA":
            ini[k] = fin[k] = fecha_firma
            return
        if preds == ["-FIRMA"]:                          # actividad previa a la firma
            ini[k] = fecha_firma - dt.timedelta(days=d)
            fin[k] = fecha_firma - dt.timedelta(days=1)
            return
        listos = []
        for p in preds:
            calc(p)
            listos.append(ini[p] if p in HITO_INICIO else fin[p] + dt.timedelta(days=1))
        listo = max(listos)
        if d == 0:                                       # hito
            ini[k] = fin[k] = listo if k in HITO_INICIO else listo - dt.timedelta(days=1)
        else:
            ini[k], fin[k] = listo, listo + dt.timedelta(days=d - 1)

    for a in ACT:
        calc(a[1])

    otros = sum(a[5] for a in ACT if a[5])
    filas = []
    def fase(k):
        if k in FASE_PREVIA: return "Previa"
        if k in FASE_CIERRE: return "Cierre"
        return "Implementación" if fin[k] <= fin["ACTA"] else "Operación"

    filas.append(dict(nivel=1, codigo="1", nombre=PROYECTO, tipo="Proyecto"))
    for ec, en, rama in ENTREGABLES:
        filas.append(dict(nivel=2, codigo=ec, nombre=en, tipo="Entregable", alcance=rama))
        for pc, pn, *_ in [p for p in PAQUETES if p[0].startswith(ec + ".")]:
            filas.append(dict(nivel=3, codigo=pc, nombre=pn, tipo="Paquete de trabajo", alcance=rama))
            for n, a in enumerate([a for a in ACT if a[0] == pc], 1):
                paq, k, nom, d, preds, costo, resp, base = a
                filas.append(dict(nivel=4, codigo=f"{pc}.{n}", nombre=nom, clave=k, alcance=rama,
                                  tipo="Hito" if _dur(d, plazo_servicio) == 0 else "Actividad",
                                  dur=_dur(d, plazo_servicio), ini=ini[k], fin=fin[k], preds=preds,
                                  costo=(monto - otros) if costo is None else costo, resp=resp, base=base, fase=fase(k)))
    for i, f in enumerate(filas, 1):
        f["id"] = i
    idx = {f["clave"]: f["id"] for f in filas if "clave" in f}
    # resúmenes (de abajo hacia arriba)
    for f in reversed(filas):
        if f["nivel"] < 4:
            hijos = [h for h in filas if h["nivel"] == f["nivel"] + 1 and (f["nivel"] == 1 or h["codigo"].rsplit(".", 1)[0] == f["codigo"])]
            f["ini"] = min(h["ini"] for h in hijos)
            f["fin"] = max(h["fin"] for h in hijos)
            f["dur"] = (f["fin"] - f["ini"]).days + 1
            f["costo"] = sum(h["costo"] for h in hijos)

    dias_act = (fin["ACTA"] - fecha_firma).days
    tabla = pd.DataFrame([{
        "Id": f["id"], "EDT": f["codigo"], "Tarea": "   " * (f["nivel"] - 1) + f["nombre"], "Tipo": f["tipo"],
        "Duración (d)": f["dur"], "Comienzo": f["ini"].strftime("%d/%m/%Y"), "Fin": f["fin"].strftime("%d/%m/%Y"),
        "Predecesoras": ";".join(str(idx[p]) for p in f.get("preds", []) if p in idx), "Responsable": f.get("resp", ""),
        "Costo (S/)": round(f["costo"], 2)} for f in filas])

    return {
        "filas": filas, "idx": idx, "tabla": tabla, "fecha_firma": fecha_firma, "plazo_activacion": plazo_activacion,
        "monto": monto, "plazo_servicio": plazo_servicio, "lineas": lineas,
        "cumple_activacion": dias_act <= plazo_activacion,
        "metricas": {
            "Firma del contrato": fecha_firma.strftime("%d/%m/%Y"),
            "Acta de implementación": fin["ACTA"].strftime("%d/%m/%Y"),
            "Días de activación usados": f"{dias_act} de {plazo_activacion}",
            "Inicio del servicio": ini["INICIO"].strftime("%d/%m/%Y"),
            "Fin del servicio": fin["OPER"].strftime("%d/%m/%Y"),
            "Cierre del contrato": fin["CIERRE"].strftime("%d/%m/%Y"),
            "Actividades": sum(1 for f in filas if f["tipo"] == "Actividad"),
            "Hitos": sum(1 for f in filas if f["tipo"] == "Hito"),
            "Costo total (S/)": round(filas[0]["costo"], 2),
        },
        "ruta_critica": "Portabilidad numérica: 5 días de SUNAFIL + 5 de trámite + 4 de activación de líneas "
                        f"= {(fin['LIN2'] - fecha_firma).days} de los {plazo_activacion} días de activación.",
    }


# ----------------------------------------------------------------------------------------
# 3. EXCEL CON FÓRMULAS
# ----------------------------------------------------------------------------------------
_F = "Calibri"
_AZUL, _AZUL2, _HIELO, _ORO = "1E2761", "2E3E7E", "DCE6F7", "E8B54D"
_l = Side(style="thin", color="B7BCC8")
_B = Border(left=_l, right=_l, top=_l, bottom=_l)
_FD, _FM = "dd/mm/yyyy", '"S/ "#,##0.00'
_fill = lambda c: PatternFill("solid", start_color=c)


def _enc(ws, fila, cols, anchos):
    for i, (c, w) in enumerate(zip(cols, anchos or [None] * len(cols)), 1):
        x = ws.cell(row=fila, column=i, value=c)
        x.font = Font(name=_F, bold=True, color="FFFFFF", size=10)
        x.fill = _fill(_AZUL2)
        x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        x.border = _B
        if w:
            ws.column_dimensions[L(i)].width = w


def _titulo(ws, texto, ncol):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    ws["A1"] = texto
    ws["A1"].font = Font(name=_F, size=14, bold=True, color="FFFFFF")
    ws["A1"].fill = _fill(_AZUL)
    ws["A2"] = SUBTITULO
    ws["A2"].font = Font(name=_F, size=9, italic=True, color="5A6070")


def excel_cronograma(cr):
    filas, idx = cr["filas"], cr["idx"]
    wb = Workbook()

    # ---------- Resumen (parámetros arriba) ----------
    rs = wb.active
    rs.title = "Resumen"
    _titulo(rs, "Resumen del cronograma", 3)
    rs.column_dimensions["A"].width = 52
    rs.column_dimensions["B"].width = 22
    rs.column_dimensions["C"].width = 60
    par = [("Fecha de firma del contrato", dt.datetime.combine(cr["fecha_firma"], dt.time()), _FD, "Dato de entrada (editable)"),
           ("Plazo de activación (días)", cr["plazo_activacion"], "0", "Contrato, cláusula quinta (editable)"),
           ("Plazo del servicio (días)", cr["plazo_servicio"], "0", "Contrato, cláusula quinta (editable)"),
           ("Monto del contrato (S/)", cr["monto"], _FM, "Contrato, cláusula tercera (editable)"),
           ("N.° de líneas", cr["lineas"], "#,##0", "Contrato, cláusula segunda (editable)")]
    rs["A4"] = "PARÁMETROS (celdas amarillas: puedes cambiarlas y todo se recalcula)"
    rs["A4"].font = Font(name=_F, bold=True, color=_AZUL)
    for k, (a, v, nf, src) in enumerate(par, 5):
        rs.cell(row=k, column=1, value=a).font = Font(name=_F, bold=True)
        c = rs.cell(row=k, column=2, value=v)
        c.font = Font(name=_F, bold=True, color="0000FF")
        c.fill = _fill("FFFF00")
        c.number_format = nf
        rs.cell(row=k, column=3, value=src).font = Font(name=_F, size=9, italic=True, color="5A6070")
    P_FIRMA, P_ACT, P_SERV, P_MONTO, P_LIN = "Resumen!$B$5", "Resumen!$B$6", "Resumen!$B$7", "Resumen!$B$8", "Resumen!$B$9"

    # ---------- Cronograma ----------
    ws = wb.create_sheet("Cronograma")
    _titulo(ws, f"Cronograma — {PROYECTO} (vista tipo MS Project)", 13)
    cols = ["Id", "EDT", "Nombre de tarea", "Tipo", "Fase", "Duración (días)", "Comienzo", "Fin", "Predecesoras (Id)",
            "Base del plazo", "Responsable", "Costo (S/)", "Nivel"]
    _enc(ws, 4, cols, [5, 10, 58, 16, 14, 10, 12, 12, 13, 24, 24, 15, 6])
    ws.freeze_panes = "D5"
    R0 = 5
    row = {f["id"]: R0 + f["id"] - 1 for f in filas}
    rk = {f["clave"]: row[f["id"]] for f in filas if "clave" in f}
    otros = [row[f["id"]] for f in filas if f["nivel"] == 4 and f["clave"] != "OPER"]
    for f in filas:
        n, lv = row[f["id"]], f["nivel"]
        vals = [f["id"], f["codigo"], "    " * (lv - 1) + f["nombre"], f["tipo"], f.get("fase", ""), None, None, None,
                ";".join(str(idx[p]) for p in f.get("preds", []) if p in idx), f.get("base", ""), f.get("resp", ""), None, lv]
        for c, v in enumerate(vals, 1):
            ws.cell(row=n, column=c, value=v)
        if lv == 4:
            k, d = f["clave"], f["dur"]
            a = next(x for x in ACT if x[1] == k)
            dur_f = P_SERV if a[3] == "SERV" else f"{P_SERV}+10" if a[3] == "SERV+10" else d
            ws.cell(row=n, column=6, value=f"={dur_f}" if isinstance(dur_f, str) else dur_f)
            if k == "FIRMA":
                ws.cell(row=n, column=7, value=f"={P_FIRMA}")
            elif a[4] == ["-FIRMA"]:
                ws.cell(row=n, column=7, value=f"={P_FIRMA}-F{n}")
            else:
                listos = ",".join(f"G{rk[p]}" if p in HITO_INICIO else f"H{rk[p]}+1" for p in a[4])
                resta = "" if (d > 0 or k in HITO_INICIO) else "-1"
                ws.cell(row=n, column=7, value=f"=MAX({listos}){resta}")
            ws.cell(row=n, column=8, value=f"=G{n}" if d == 0 else f"=G{n}+F{n}-1")
            if k == "OPER":
                ws.cell(row=n, column=12, value=f"={P_MONTO}-(" + "+".join(f"L{x}" for x in otros) + ")")
            else:
                ws.cell(row=n, column=12, value=f["costo"]).font = Font(name=_F, size=10, color="0000FF")
        else:
            hijos = [row[h["id"]] for h in filas if h["nivel"] == lv + 1 and (lv == 1 or h["codigo"].rsplit(".", 1)[0] == f["codigo"])]
            ws.cell(row=n, column=7, value="=MIN(" + ",".join(f"G{h}" for h in hijos) + ")")
            ws.cell(row=n, column=8, value="=MAX(" + ",".join(f"H{h}" for h in hijos) + ")")
            ws.cell(row=n, column=6, value=f"=H{n}-G{n}+1")
            ws.cell(row=n, column=12, value="=" + "+".join(f"L{h}" for h in hijos))
        fondo = {1: _ORO, 2: _HIELO, 3: "EEF1F8"}.get(lv)
        for c in range(1, 14):
            x = ws.cell(row=n, column=c)
            x.border = _B
            if not (x.font and x.font.color and x.font.color.rgb == "000000FF"):
                x.font = Font(name=_F, size=10, bold=lv <= 3, italic=f["tipo"] == "Hito",
                              color="7A4B00" if f["tipo"] == "Hito" else "000000")
            if fondo:
                x.fill = _fill(fondo)
            x.alignment = Alignment(vertical="center", wrap_text=c in (3, 10, 11),
                                    horizontal="center" if c in (1, 2, 4, 5, 6, 7, 8, 9, 13) else None)
        if lv == 4 and f["base"] == "Plazo del TDR o del contrato":
            ws.cell(row=n, column=10).fill = _fill("E2F0D9")
        elif lv == 4:
            ws.cell(row=n, column=10).fill = _fill("FFF2CC")
        ws.cell(row=n, column=7).number_format = _FD
        ws.cell(row=n, column=8).number_format = _FD
        ws.cell(row=n, column=12).number_format = _FM
    LAST = R0 + len(filas) - 1
    nota = LAST + 2
    for t in ["Las actividades (verbos) abren cada paquete de trabajo de la EDT, como en MS Project: 1.5.1 → 1.5.1.1, 1.5.1.2…",
              "Base del plazo: verde = plazo del TDR o del contrato; amarillo = estimación del equipo.",
              "Costos en azul = estimación referencial del grupo. La operación del servicio recibe el saldo para que el total sea igual al monto del contrato (suma alzada)."]:
        ws.cell(row=nota, column=3, value="• " + t).font = Font(name=_F, size=9, italic=True, color="5A6070")
        nota += 1

    # ---------- Gantt implementación (diario) ----------
    g = wb.create_sheet("Gantt implementación")
    ventana = max(cr["plazo_activacion"] + 16, 30)
    g["A1"] = f"Ventana de activación: {cr['plazo_activacion']} días calendario desde el día siguiente a la firma"
    g["A1"].font = Font(name=_F, size=13, bold=True, color=_AZUL)
    g["A2"] = "Rojo = hito · Dorado = actividad · Línea roja vertical = límite del plazo de activación"
    g["A2"].font = Font(name=_F, size=9, italic=True, color="5A6070")
    _enc(g, 4, ["EDT", "Actividad", "Comienzo", "Fin"], [10, 50, 11, 11])
    for d in range(ventana):
        c = g.cell(row=4, column=5 + d, value=f"={P_FIRMA}+{d - 5}")
        c.number_format = "dd/mm"
        c.font = Font(name=_F, size=7, color="FFFFFF")
        c.fill = _fill(_AZUL2)
        c.alignment = Alignment(text_rotation=90, horizontal="center")
        g.column_dimensions[L(5 + d)].width = 2.8
    g.row_dimensions[4].height = 40
    fin_vent = cr["fecha_firma"] + dt.timedelta(days=ventana - 6)
    gr = 5
    for f in filas:
        if f["nivel"] != 4 or f["ini"] > fin_vent:
            continue
        n = row[f["id"]]
        for c, v in enumerate([f"=Cronograma!B{n}", f"=Cronograma!C{n}", f"=Cronograma!G{n}", f"=Cronograma!H{n}"], 1):
            x = g.cell(row=gr, column=c, value=v)
            x.font = Font(name=_F, size=9)
            if c >= 3:
                x.number_format = _FD
        gr += 1
    rng = f"E5:{L(4 + ventana)}{gr - 1}"
    g.conditional_formatting.add(rng, FormulaRule(formula=["AND(ISNUMBER(SEARCH(\"Hito:\",$B5)),E$4=$C5)"], fill=_fill("C00000")))
    g.conditional_formatting.add(rng, FormulaRule(formula=["AND(E$4>=$C5,E$4<=$D5)"], fill=_fill(_ORO)))
    g.conditional_formatting.add(rng, FormulaRule(formula=[f"E$4={P_FIRMA}+{P_ACT}"],
                                                  border=Border(right=Side(style="thick", color="C00000"))))
    g.freeze_panes = "E5"

    # ---------- Gantt general (mensual) ----------
    m = wb.create_sheet("Gantt general")
    m["A1"] = "Gantt general — entregables y paquetes de trabajo por mes"
    m["A1"].font = Font(name=_F, size=13, bold=True, color=_AZUL)
    _enc(m, 4, ["EDT", "Entregable / paquete", "Comienzo", "Fin"], [8, 40, 11, 11])
    meses = cr["plazo_servicio"] // 30 + 3
    for k in range(meses):
        c = m.cell(row=4, column=5 + k, value=f"=DATE(YEAR({P_FIRMA}),MONTH({P_FIRMA})+{k},1)")
        c.number_format = "mmm-yy"
        c.font = Font(name=_F, size=8, color="FFFFFF")
        c.fill = _fill(_AZUL2)
        c.alignment = Alignment(text_rotation=90, horizontal="center")
        m.column_dimensions[L(5 + k)].width = 4
    m.row_dimensions[4].height = 42
    mr = 5
    for f in filas:
        if f["nivel"] not in (2, 3):
            continue
        n = row[f["id"]]
        for c, v in enumerate([f"=Cronograma!B{n}", f"=Cronograma!C{n}", f"=Cronograma!G{n}", f"=Cronograma!H{n}"], 1):
            x = m.cell(row=mr, column=c, value=v)
            x.font = Font(name=_F, size=9, bold=f["nivel"] == 2)
            if c >= 3:
                x.number_format = _FD
            if f["nivel"] == 2:
                x.fill = _fill(_HIELO)
        mr += 1
    rng = f"E5:{L(4 + meses)}{mr - 1}"
    m.conditional_formatting.add(rng, FormulaRule(formula=["AND(E$4<=$D5,DATE(YEAR(E$4),MONTH(E$4)+1,0)>=$C5,LEN($A5)<=3)"], fill=_fill(_AZUL)))
    m.conditional_formatting.add(rng, FormulaRule(formula=["AND(E$4<=$D5,DATE(YEAR(E$4),MONTH(E$4)+1,0)>=$C5)"], fill=_fill(_ORO)))
    m.freeze_panes = "E5"

    # ---------- Costos ----------
    c = wb.create_sheet("Costos")
    _titulo(c, "Costos del proyecto", 5)
    for col, w in zip("ABCDE", [10, 40, 16, 18, 50]):
        c.column_dimensions[col].width = w
    _enc(c, 4, ["", "Concepto", "", "Monto", "Fuente"], None)
    ind = [("Monto contractual total", f"={P_MONTO}", "Contrato, cláusula tercera"),
           ("Costo mensual del servicio", f"={P_MONTO}/ROUND({P_SERV}/30.4,0)", "Monto entre los meses del servicio (24)"),
           ("Costo por línea al mes", f"=D6/{P_LIN}", "Costo mensual entre el N.° de líneas"),
           ("Garantía de fiel cumplimiento (10%)", f"={P_MONTO}*10%", "Contrato, cláusula séptima"),
           ("Penalidad máxima por mora (10%)", f"={P_MONTO}*10%", "Bases, penalidades (tope del 10%)")]
    for k, (a, f_, s) in enumerate(ind, 5):
        c.cell(row=k, column=2, value=a).font = Font(name=_F, bold=True)
        x = c.cell(row=k, column=4, value=f_)
        x.number_format = _FM
        c.cell(row=k, column=5, value=s).font = Font(name=_F, size=9, italic=True, color="5A6070")
    _enc(c, 11, ["EDT", "Entregable", "Alcance", "Costo (S/)", "% del total"], None)
    ent_rows = {f["codigo"]: row[f["id"]] for f in filas if f["nivel"] == 2}
    r0 = 12
    for k, (ec, en, rama) in enumerate(ENTREGABLES):
        rr = r0 + k
        c.cell(row=rr, column=1, value=ec)
        c.cell(row=rr, column=2, value=en)
        c.cell(row=rr, column=3, value=rama)
        x = c.cell(row=rr, column=4, value=f"=Cronograma!L{ent_rows[ec]}")
        x.number_format = _FM
        x.font = Font(name=_F, color="008000")
        c.cell(row=rr, column=5, value=f"=IF($D$21=0,0,D{rr}/$D$21)").number_format = "0.0%"
    c.cell(row=21, column=2, value="TOTAL").font = Font(name=_F, bold=True)
    c.cell(row=21, column=4, value="=SUM(D12:D20)").number_format = _FM
    c.cell(row=21, column=5, value="=SUM(E12:E20)").number_format = "0.0%"
    c.cell(row=22, column=2, value="Control: total vs. monto del contrato")
    c.cell(row=22, column=4, value=f'=IF(ROUND(D21-{P_MONTO},2)=0,"CUADRA","NO CUADRA")').font = Font(name=_F, bold=True, color="2E8B57")
    c.cell(row=24, column=2, value="Nota: el contrato es a suma alzada y no desagrega el precio por entregable; la distribución "
                                   "es referencial (costos en azul en la hoja Cronograma).").font = Font(name=_F, size=9, italic=True)

    # ---------- Resumen (resultados con fórmulas) ----------
    rs["A11"] = "RESULTADOS"
    rs["A11"].font = Font(name=_F, bold=True, color=_AZUL)
    tipos = f"Cronograma!$D$5:$D${LAST}"
    res = [("Acta de culminación de implementación", f"=Cronograma!G{rk['ACTA']}", _FD),
           ("Días de activación usados", f"=Cronograma!G{rk['ACTA']}-{P_FIRMA}", "0"),
           ("¿Cumple el plazo de activación?", f'=IF(B13<={P_ACT},"CUMPLE","NO CUMPLE")', None),
           ("Inicio del servicio", f"=Cronograma!G{rk['INICIO']}", _FD),
           ("Fin del servicio", f"=Cronograma!H{rk['OPER']}", _FD),
           ("Cierre del contrato", f"=Cronograma!G{rk['CIERRE']}", _FD),
           ("Duración total del proyecto (días)", f"=Cronograma!F{R0}", "0"),
           ("Entregables de 1er nivel", f'=COUNTIF({tipos},"Entregable")', "0"),
           ("Paquetes de trabajo", f'=COUNTIF({tipos},"Paquete de trabajo")', "0"),
           ("Actividades", f'=COUNTIF({tipos},"Actividad")', "0"),
           ("Hitos", f'=COUNTIF({tipos},"Hito")', "0"),
           ("Actividades con plazo del TDR o del contrato", f'=COUNTIF(Cronograma!$J$5:$J${LAST},"Plazo del TDR o del contrato")-COUNTIF({tipos},"Hito")', "0"),
           ("Costo total del proyecto", f"=Cronograma!L{R0}", _FM),
           ("Control de costo vs. contrato", "=Costos!D22", None),
           ("Ruta crítica de la implementación", cr["ruta_critica"], None)]
    for k, (a, f_, nf) in enumerate(res, 12):
        rs.cell(row=k, column=1, value=a).font = Font(name=_F, bold=True)
        x = rs.cell(row=k, column=2, value=f_)
        x.font = Font(name=_F, bold=True, color="008000" if str(f_).startswith("=") else "000000")
        x.alignment = Alignment(horizontal="left", wrap_text=k == 26)
        if nf:
            x.number_format = nf
    rs.merge_cells("B26:C26")
    rs.row_dimensions[26].height = 32

    for hoja in wb:
        hoja.page_setup.orientation = "landscape"
        hoja.page_setup.fitToWidth, hoja.page_setup.fitToHeight = 1, 0
        hoja.sheet_properties.pageSetUpPr.fitToPage = True
    buf = io.BytesIO()
    wb.calculation.fullCalcOnLoad = True
    wb.save(buf)
    return buf.getvalue()
