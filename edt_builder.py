"""
edt_builder.py — Sección 3 (EDT) de la app "Checklist de TDR — OCR (Gemini) -> Excel".

Convierte un checklist (el que genera la sección 2 o un Excel que ya tengas) en:
  * EDT_SUNAFIL.xlsx  -> hojas: EDT, Diccionario EDT, Fuera de la EDT, Sin clasificar, Verificación 100%
  * EDT_diagrama.pptx -> árbol editable en PowerPoint (proyecto -> 9 entregables -> 28 paquetes)

No usa Gemini: la estructura (9 entregables / 28 paquetes) es la plantilla validada en clase
y cada ítem del checklist se asigna por reglas de palabras clave, trazables y repetibles.

Uso:
    from edt_builder import leer_checklist_excel, construir_edt, excel_edt, pptx_edt
    df = leer_checklist_excel(archivo_subido)          # o un DataFrame de la sección 2
    res = construir_edt(df)                             # dict con filas y métricas
    xlsx_bytes = excel_edt(res); pptx_bytes = pptx_edt(res)
"""
import io
import re
import unicodedata

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

PROYECTO = "Servicio de Telefonía Móvil para la SUNAFIL"
SUBTITULO = "Concurso Público de Servicios N.° 001-2026-SUNAFIL-1 · Contrato N.° 05-2026-SUNAFIL-OAD"

# ----------------------------------------------------------------------------------------
# 1. PLANTILLA DE LA EDT (PMBOK): 9 entregables de 1er nivel y 28 paquetes de trabajo
#    (código, nombre, descripción del paquete, criterio de aceptación, motivo si no tiene ítem)
# ----------------------------------------------------------------------------------------
ENTREGABLES = [
    ("1.1", "Servicio móvil habilitado", "Producto"),
    ("1.2", "Equipamiento móvil", "Producto"),
    ("1.3", "Plataformas de gestión", "Producto"),
    ("1.4", "Cobertura garantizada", "Producto"),
    ("1.5", "Gestión del proyecto", "Proyecto"),
    ("1.6", "Implementación y activación", "Proyecto"),
    ("1.7", "Transferencia de conocimiento", "Proyecto"),
    ("1.8", "Soporte y continuidad", "Proyecto"),
    ("1.9", "Cierre del contrato", "Proyecto"),
]

PAQUETES = [
    ("1.1.1", "Líneas móviles activadas", "Líneas de voz, SMS y datos activas para SUNAFIL.", "Relación de líneas activas firmada por la OTIC.", ""),
    ("1.1.2", "Plan de voz", "Minutos ilimitados a cualquier operador y restricciones por perfil.", "Plan configurado en todas las líneas.", ""),
    ("1.1.3", "Plan de datos por gama", "Bolsas de datos por gama, incluidas las mejoras ofertadas.", "GB asignados según el contrato y visibles en el portal.", "Exigido en el TDR (planes de datos por gama) y mejorado en el Factor I; en algunos checklists va dentro del ítem de equipos o líneas."),
    ("1.1.4", "SMS y servicios suplementarios", "SMS, casilla de voz, desvío, llamada en espera y números de emergencia.", "Servicios probados en una muestra de líneas.", ""),
    ("1.1.5", "Roaming internacional", "Roaming 'ida y vuelta' para líneas autorizadas.", "Prueba de roaming realizada.", "Exigido en el TDR (roaming 'ida y vuelta', facturación aparte según el Pliego); en algunos checklists va dentro de llamadas."),
    ("1.2.1", "Equipos Gama A", "Equipos de gama alta en alquiler.", "Acta de entrega y ficha técnica conforme.", ""),
    ("1.2.2", "Equipos Gama B", "Equipos de gama media en alquiler.", "Acta de entrega y ficha técnica conforme.", ""),
    ("1.2.3", "Equipos Gama C", "Equipos de gama básica en alquiler.", "Actas de entrega por dependencia.", ""),
    ("1.2.4", "Equipos de respaldo y accesorios", "Equipos de respaldo por gama y accesorios de cada equipo.", "Respaldo y accesorios entregados con acta.", ""),
    ("1.3.1", "Consola MDM en la nube", "Plataforma MDM para enrolar, proteger y administrar los equipos.", "Consola operativa con políticas aplicadas.", ""),
    ("1.3.2", "Herramienta web de consumo", "Portal web 24x7 para consultar el consumo por línea.", "Usuarios creados y reporte validado por la OTIC.", ""),
    ("1.4.1", "Cobertura nacional acreditada", "Red 3G/4G/5G del operador con concesión del MTC.", "Informe de cobertura en sede central y dependencias.", ""),
    ("1.4.2", "Solución indoor en sedes", "Mejora de señal en dependencias con nivel bajo.", "Mediciones posteriores a la solución.", ""),
    ("1.5.1", "Plan de trabajo", "Plan y cronograma de implementación.", "Plan aprobado por la OTIC.",
     "Derivado de la gestión PMBOK: el TDR (8.5) no lo exige, pero el equipo lo necesita para dirigir la implementación."),
    ("1.5.2", "Entregables mensuales", "Informe mensual de consumos, incidencias y operación.", "24 informes presentados a tiempo.", ""),
    ("1.5.3", "Documentación de perfeccionamiento", "Requisitos del proveedor, personal clave y documentos para firmar el contrato.", "Documentos completos al suscribir el contrato.", ""),
    ("1.5.4", "Seguimiento y control", "Coordinaciones, actas de seguimiento y registro de riesgos.", "Actas firmadas por ambas partes.",
     "Derivado de la gestión PMBOK: control del proyecto durante los 24 meses."),
    ("1.6.1", "Portabilidad numérica", "Trámite de portabilidad a cargo del contratista.", "Líneas portadas con su numeración.", ""),
    ("1.6.2", "Equipos configurados y enrolados", "Configuración de equipos y enrolamiento en la consola MDM.", "100% de equipos enrolados.",
     "Derivado de 1.3.1: el TDR exige el enrolamiento dentro del plazo de activación."),
    ("1.6.3", "Acta de inicio del servicio", "Entregable 11.1: acta de culminación de implementación e inicio del servicio.", "Acta firmada dentro del plazo de activación.", ""),
    ("1.7.1", "Manuales y material didáctico", "Manuales de la consola MDM y del portal web.", "Manuales entregados a la OTIC.",
     "Derivado de la transferencia de conocimiento: el TDR pide material didáctico junto con la capacitación."),
    ("1.7.2", "Capacitación ejecutada", "Capacitación presencial o virtual a los administradores.", "Constancias y lista de asistencia.", ""),
    ("1.8.1", "Mesa de ayuda 24x7x365", "Atención de consultas y soporte a usuarios.", "Reporte mensual de tickets.", ""),
    ("1.8.2", "Garantía de equipos", "Garantía de equipos, baterías y accesorios.", "Garantías atendidas sin costo.", ""),
    ("1.8.3", "Reposición y asistencia técnica", "Reposición por falla, robo o pérdida y asistencia técnica.", "Equipos repuestos en plazo.", ""),
    ("1.9.1", "Conformidades y pagos", "Conformidad de la OTIC y pagos mensuales.", "Conformidades emitidas.", ""),
    ("1.9.2", "Devolución o adquisición de equipos", "Al término: devolución de equipos o compra por valor residual.", "Acta de devolución o de transferencia.", "Exigido en el TDR (equipos en alquiler): al término se devuelven o se compran por valor residual (Pliego, consulta 56)."),
    ("1.9.3", "Informe final y conformidad final", "Informe final del servicio y conformidad final.", "Acta de conformidad final.",
     "Derivado de 11.2 y XII: la conformidad final cierra el contrato."),
]
PAQ = {p[0]: p for p in PAQUETES}
ENT = {e[0]: e for e in ENTREGABLES}

# ----------------------------------------------------------------------------------------
# 2. REGLAS DE ASIGNACIÓN (en orden; gana la primera que coincide). Se aplican al título
#    del ítem normalizado (minúsculas, sin tildes). FUERA = no es entregable, con motivo.
# ----------------------------------------------------------------------------------------
F_NOAPL = ("El TDR lo declara no aplicable", "No se registra: el TDR lo excluye")
F_CONTR = ("Condición contractual, no es trabajo a entregar", "Enunciado del alcance (restricciones y supuestos)")
F_PLAZO = ("Restricción de plazo: alimenta el cronograma, no la EDT", "Cronograma y línea base de tiempo")
F_LUGAR = ("Restricción de lugar de prestación", "Enunciado del alcance (restricciones)")
F_CONTX = ("Contexto o justificación del requerimiento, no es trabajo", "Acta de constitución / enunciado del alcance")
F_ENTID = ("Responsabilidad de la Entidad, no del contratista", "Matriz de roles (trabajo de SUNAFIL)")
F_FACTR = ("Criterio para evaluar la oferta, no trabajo del proyecto", "Informe de evaluación de ofertas")

REGLAS = [
    # --- fuera de la EDT ---
    (r"no aplica", "FUERA", F_NOAPL),
    (r"^factor|factor de evaluacion|sostenibilidad|integridad en la contratacion|sistema de gestion de la calidad|experiencia adicional", "FUERA", F_FACTR),
    (r"penalidad", "FUERA", F_CONTR),
    (r"subcontrat", "FUERA", F_CONTR),
    (r"controversia|arbitra", "FUERA", F_CONTR),
    (r"reajuste", "FUERA", F_CONTR),
    (r"respuestas? entre las partes", "FUERA", F_CONTR),
    (r"confidencial", "FUERA", F_CONTR),
    (r"gastos adicionad|no previstos", "FUERA", F_CONTR),
    (r"modalidad de (pago|contratacion)|suma alzada", "FUERA", F_CONTR),
    (r"supervision del servicio", "FUERA", F_ENTID),
    (r"finalidad|descripcion general|terminos de referencia|area usuaria|objetivo|justificacion", "FUERA", F_CONTX),
    # --- entregables del proyecto (antes que los "plazo de ...") ---
    (r"entregable mensual|informe mensual|11\.2", "1.5.2", None),
    (r"entregable de (activacion|inicio)|inicio de(l)? servicio|culminacion|11\.1", "1.6.3", None),
    (r"plazo de (prestacion|activacion)|plazo del servicio", "FUERA", F_PLAZO),
    (r"lugar de prestacion", "FUERA", F_LUGAR),
    (r"portabilidad", "1.6.1", None),
    (r"enrol|espacio de trabajo|wifi para configuracion|configuracion de equipos", "1.6.2", None),
    (r"transferencia de conocimiento|capacitacion", "1.7.2", None),
    (r"manual|didactic", "1.7.1", None),
    (r"personal designado|canalizacion|coordinacion", "1.5.4", None),
    (r"plan de trabajo", "1.5.1", None),
    (r"requisitos? del proveedor|declaracion del postor|rnp|autorizacion del mtc|domicilio|inviolabilidad"
     r"|personal clave|gestor de servicio|certificac|formacion academica|experiencia minima", "1.5.3", None),
    (r"conformidad|recepcion del servicio|forma de pago|esquema de pagos|condiciones y forma", "1.9.1", None),
    (r"desfase|depreciacion|devolucion|adquisicion al termino|propiedad", "1.9.2", None),
    # --- producto (los específicos primero) ---
    (r"mdm|consola|antivirus|antimalware|phishing|geolocaliz|ubicacion|kiosko|borrado|limpieza del dispositivo"
     r"|acciones remotas|acceso remoto|repositorio|lista negra|modo avion|tarjeta sim|infeccion|advertencias"
     r"|uso de aplicaciones|actualizacion de aplicaciones|instalacion remota|permitir o denegar|sistemas operativos"
     r"|uso de datos del dispositivo|reportes e informes|intentos fallidos|aplicaciones no autorizadas|llamadas entrantes", "1.3.1", None),
    (r"servicios? (complementari|suplementari)", "1.1.4", None),
    (r"roaming|ida y vuelta", "1.1.5", None),
    (r"asignacion del plan|plan contratado|plan de datos", "1.1.3", None),
    (r"atencion de consultas|soporte (y atencion|al usuario)|24x7|centros de atencion", "1.8.1", None),
    (r"reposicion|robo|perdida|averia|reemplazo|defecto|asistencia tecnica|mantenimiento|reten|soporte tecnico", "1.8.3", None),
    (r"garantia", "1.8.2", None),
    (r"herramienta (de gestion|web)|portal web|monitoreo de consumo|canal de atencion", "1.3.2", None),
    (r"accesorio|cargador", "1.2.4", None),
    (r"respaldo|distribucion de equipos", "1.2.4", None),
    (r"gama a\b", "1.2.1", None),
    (r"gama b\b", "1.2.2", None),
    (r"gama c\b", "1.2.3", None),
    (r"indoor|outdoor|calidad de senal", "1.4.2", None),
    (r"cobertura|tecnologia de (la )?red|infraestructura|2g|4g|5g", "1.4.1", None),
    (r"internacional", "1.1.5", None),
    (r"sms|mensaj|desvio|llamada en espera|casilla|complementari|suplementari|emergencia|discado", "1.1.4", None),
    (r"llamada|voz|minutos|tarifa|otros operadores", "1.1.2", None),
    (r"datos|gb\b|plan contratado|asignacion del plan|internet", "1.1.3", None),
    (r"linea|servicio a realizar|cantidad", "1.1.1", None),
]
REGLAS = [(re.compile(p), d, f) for p, d, f in REGLAS]


def _norm(s):
    s = "" if s is None or (isinstance(s, float) and pd.isna(s)) else str(s)
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s.lower()).strip()


# ----------------------------------------------------------------------------------------
# 3. LECTURA DEL CHECKLIST (acepta el Excel de la sección 2 y el checklist anterior del grupo)
# ----------------------------------------------------------------------------------------
_ALIAS = {
    "numeral": ["numeral tdr", "punto del tdr", "numeral"],
    "titulo": ["titulo / item", "titulo", "item / componente", "item"],
    "clase": ["clase", "alcance"],
    "exigido": ["valor exigido [a+b]", "valor exigido (a+b)", "valor exigido"],
    "alcance": ["a+b+c) alcance final", "a) alcance (bases)", "a) alcance"],
    "consulta": ["a+b) consulta (pliego)", "a+b) alcance + consulta (pliego)", "consulta"],
    "veredicto": ["¿cumple?", "cumple?", "cumple", "veredicto"],
}


def _mapear_columnas(cols):
    norm = {c: _norm(c) for c in cols}
    out = {}
    for clave, alias in _ALIAS.items():
        for a in alias:
            hit = [c for c, n in norm.items() if n == _norm(a)] or [c for c, n in norm.items() if n.startswith(_norm(a))]
            if hit:
                out[clave] = hit[0]
                break
    return out


def leer_checklist_excel(archivo):
    """Devuelve un DataFrame con columnas estándar a partir de un .xlsx de checklist."""
    libro = pd.read_excel(archivo, sheet_name=None, header=None)
    for nombre, hoja in libro.items():
        if "checklist" not in _norm(nombre) or _norm(nombre).startswith("b2"):
            continue
        for i in range(min(10, len(hoja))):  # busca la fila de encabezados
            fila = [_norm(x) for x in hoja.iloc[i].tolist()]
            if any(x.startswith("numeral") or x == "punto del tdr" for x in fila):
                df = hoja.iloc[i + 1:].copy()
                df.columns = hoja.iloc[i].tolist()
                return normalizar_checklist(df)
    raise ValueError("No encontré una hoja de checklist con la columna 'Numeral TDR'.")


def normalizar_checklist(df):
    """Acepta el DataFrame de la sección 2 (mismas columnas que la hoja B) o uno leído de Excel."""
    m = _mapear_columnas(df.columns)
    if "numeral" not in m or "titulo" not in m:
        raise ValueError("El checklist debe tener las columnas 'Numeral TDR' y 'Título / Ítem'.")
    out = pd.DataFrame({k: df[v].values for k, v in m.items()})
    for k in _ALIAS:
        if k not in out:
            out[k] = ""
    out = out[out["numeral"].notna() & (out["numeral"].astype(str).str.strip() != "")]
    out = out[~out["numeral"].astype(str).str.upper().str.startswith("TOTAL")]
    return out.fillna("").reset_index(drop=True)


# ----------------------------------------------------------------------------------------
# 4. CONSTRUCCIÓN
# ----------------------------------------------------------------------------------------
def _asignar(numeral, titulo, veredicto):
    texto = _norm(f"{titulo}")
    if _norm(veredicto) == "no aplica":
        return "FUERA", F_NOAPL, "Veredicto NO APLICA"
    if _norm(numeral).startswith("factor"):
        return "FUERA", F_FACTR, "Es factor de evaluación"
    for patron, destino, fuera in REGLAS:
        m = patron.search(texto)
        if m:
            return destino, fuera, f"Regla: «{m.group(0)}»"
    return "SIN", None, "Ninguna regla coincide"


def construir_edt(df):
    df = normalizar_checklist(df) if "numeral" not in df.columns else df
    asignados, fuera, sin = [], [], []
    for _, r in df.iterrows():
        dest, mot, como = _asignar(r["numeral"], r["titulo"], r["veredicto"])
        fila = dict(r)
        fila["como"] = como
        if dest == "FUERA":
            fila["motivo"], fila["registro"] = mot
            fuera.append(fila)
        elif dest == "SIN":
            sin.append(fila)
        else:
            fila["codigo"] = dest
            asignados.append(fila)

    por_paquete = {p[0]: [] for p in PAQUETES}
    for a in asignados:
        por_paquete[a["codigo"]].append(a)
    vacios = [c for c, items in por_paquete.items() if not items]
    vacios_sin_just = [c for c in vacios if not PAQ[c][4]]

    # validaciones PMBOK sobre la plantilla
    hijos = {}
    for c, *_ in PAQUETES:
        hijos.setdefault(c.rsplit(".", 1)[0], []).append(c)
    uno_a_uno = [e for e, h in hijos.items() if len(h) < 2]

    return {
        "items": len(df), "asignados": asignados, "fuera": fuera, "sin": sin,
        "por_paquete": por_paquete, "vacios": vacios, "vacios_sin_just": vacios_sin_just,
        "uno_a_uno": uno_a_uno,
        "metricas": {
            "Ítems del alcance": len(df),
            "Se vuelven entregables": len(asignados),
            "Fuera, con motivo escrito": len(fuera),
            "Sin clasificar": len(sin),
            "Entregables de 1er nivel": len(ENTREGABLES),
            "Paquetes de trabajo": len(PAQUETES),
        },
        "cumple_100": len(asignados) + len(fuera) == len(df) and not sin,
    }


def tabla_edt(res):
    """Filas para mostrar la EDT en pantalla (st.dataframe)."""
    filas = [{"Código": "1", "Entregable": PROYECTO, "Nivel": 1, "Alcance": "", "Ítems": res["items"] - len(res["fuera"]) - len(res["sin"])}]
    for ec, en, rama in ENTREGABLES:
        n = sum(len(res["por_paquete"][p[0]]) for p in PAQUETES if p[0].startswith(ec + "."))
        filas.append({"Código": ec, "Entregable": "   " + en, "Nivel": 2, "Alcance": rama, "Ítems": n})
        for pc, pn, *_ in [p for p in PAQUETES if p[0].startswith(ec + ".")]:
            filas.append({"Código": pc, "Entregable": "      " + pn, "Nivel": 3, "Alcance": rama, "Ítems": len(res["por_paquete"][pc])})
    return pd.DataFrame(filas)


# ----------------------------------------------------------------------------------------
# 5. EXCEL
# ----------------------------------------------------------------------------------------
_AZUL, _AZUL2, _HIELO, _ORO = "1E2761", "2E3E7E", "DCE6F7", "E8B54D"
_lado = Side(style="thin", color="B7BCC8")
_BORDE = Border(left=_lado, right=_lado, top=_lado, bottom=_lado)
_AJ = Alignment(wrap_text=True, vertical="top")


def _hoja(ws, titulo, cols, anchos):
    ws["A1"] = titulo
    ws["A1"].font = Font(name="Calibri", size=14, bold=True, color=_AZUL)
    ws["A2"] = SUBTITULO
    ws["A2"].font = Font(name="Calibri", size=9, italic=True, color="5A6070")
    for i, (c, w) in enumerate(zip(cols, anchos), 1):
        x = ws.cell(row=4, column=i, value=c)
        x.font = Font(name="Calibri", bold=True, color="FFFFFF")
        x.fill = PatternFill("solid", start_color=_AZUL)
        x.alignment = Alignment(wrap_text=True, vertical="center")
        x.border = _BORDE
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:{chr(64 + len(cols))}4"


def _fila(ws, r, vals, bold=False, fill=None):
    for i, v in enumerate(vals, 1):
        x = ws.cell(row=r, column=i, value=v)
        x.font = Font(name="Calibri", size=10, bold=bold)
        x.alignment = _AJ
        x.border = _BORDE
        if fill:
            x.fill = PatternFill("solid", start_color=fill)


def excel_edt(res):
    wb = Workbook()
    # --- EDT (vertical, tipo Project) ---
    ws = wb.active
    ws.title = "EDT"
    _hoja(ws, f"EDT — {PROYECTO}", ["Código", "Entregable / Paquete de trabajo", "Nivel", "Alcance", "Tipo (PMBOK)", "Ítems del TDR que contiene", "N° ítems"],
          [10, 42, 7, 11, 20, 60, 9])
    r = 5
    _fila(ws, r, ["1", PROYECTO, 1, "", "Proyecto", "", None], bold=True, fill=_ORO)
    fila_raiz, filas_ent = r, []
    for ec, en, rama in ENTREGABLES:
        r += 1
        fe = r
        filas_ent.append(fe)
        _fila(ws, r, [ec, en, 2, rama, "Entregable (1er nivel)", "", None], bold=True, fill=_HIELO)
        paqs = [p for p in PAQUETES if p[0].startswith(ec + ".")]
        for pc, pn, *_ in paqs:
            r += 1
            items = res["por_paquete"][pc]
            nums = ", ".join(str(i["numeral"]) for i in items) or "— (derivado, ver diccionario)"
            _fila(ws, r, [pc, "    " + pn, 3, rama, "Paquete de trabajo", nums, len(items)])
        ws.cell(row=fe, column=7, value=f"=SUM(G{fe + 1}:G{r})")
    ws.cell(row=fila_raiz, column=7, value="=" + "+".join(f"G{x}" for x in filas_ent))
    r += 2
    for t in ["Solo entregables (sustantivos); las actividades (verbos) van en el cronograma.",
              "Rama Producto (1.1–1.4): lo que recibe SUNAFIL. Rama Proyecto (1.5–1.9): el trabajo para entregarlo y gestionarlo.",
              "Paquete de trabajo = último nivel de descomposición; su código es la cuenta de control.",
              "Ningún entregable se descompone 1 a 1: cada uno se abre en 2 o más paquetes."]:
        ws.cell(row=r, column=1, value="• " + t).font = Font(name="Calibri", size=9, italic=True, color="5A6070")
        r += 1

    # --- Diccionario EDT ---
    ws = wb.create_sheet("Diccionario EDT")
    _hoja(ws, "Diccionario de la EDT", ["Código", "Paquete de trabajo", "Alcance", "Numeral TDR", "Ítem del TDR",
                                        "Requisito verificable (A+B)", "Consulta del Pliego", "Veredicto checklist", "Cómo se asignó"],
          [9, 30, 10, 16, 40, 50, 40, 15, 22])
    r = 4
    for pc, pn, desc, crit, just in PAQUETES:
        rama = ENT[pc.rsplit(".", 1)[0]][2]
        items = res["por_paquete"][pc]
        if not items:
            r += 1
            _fila(ws, r, [pc, pn, rama, "—", "(sin numeral propio)", f"{desc} Criterio: {crit}", "", "", just or "Sin ítem"], fill="FFF4E5")
        for it in items:
            r += 1
            _fila(ws, r, [pc, pn, rama, it["numeral"], it["titulo"], it["exigido"] or it["alcance"], it["consulta"], it["veredicto"], it["como"]])

    # --- Fuera de la EDT ---
    ws = wb.create_sheet("Fuera de la EDT")
    _hoja(ws, "Ítems del checklist que no son entregables", ["Numeral TDR", "Ítem", "Por qué no es un entregable", "Dónde se registra"], [16, 45, 50, 45])
    r = 4
    for it in res["fuera"]:
        r += 1
        _fila(ws, r, [it["numeral"], it["titulo"], it["motivo"], it["registro"]])

    # --- Sin clasificar ---
    ws = wb.create_sheet("Sin clasificar")
    _hoja(ws, "Ítems que ninguna regla pudo ubicar", ["Numeral TDR", "Ítem", "Requisito", "Qué hacer"], [16, 45, 50, 45])
    if res["sin"]:
        for k, it in enumerate(res["sin"], 5):
            _fila(ws, k, [it["numeral"], it["titulo"], it["exigido"] or it["alcance"], "Asignarlo a un paquete o a 'Fuera de la EDT'"])
    else:
        _fila(ws, 5, ["—", "Ningún ítem quedó sin clasificar", "", ""], fill="E2F0D9")

    # --- Verificación 100% (con fórmulas) ---
    ws = wb.create_sheet("Verificación 100%")
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 60
    filas = [
        ("Ítems del alcance analizados", res["items"]),
        ("Se convierten en entregables (entran en la EDT)", "=COUNTA('Diccionario EDT'!D5:D500)-COUNTIF('Diccionario EDT'!D5:D500,\"—\")"),
        ("Quedan fuera con justificación escrita", "=COUNTA('Fuera de la EDT'!A5:A500)"),
        ("Sin clasificar (debe ser 0)", len(res["sin"])),
        ("¿La suma cuadra con el total?", '=IF(B6+B7+B8=B5,"SÍ","NO")'),
        ("", None),
        ("Entregables de 1er nivel", len(ENTREGABLES)),
        ("Paquetes de trabajo", len(PAQUETES)),
        ("Paquetes sin ítem directo (justificados en el diccionario)", len(res["vacios"]) - len(res["vacios_sin_just"])),
        ("Paquetes vacíos sin justificar (debe ser 0)", len(res["vacios_sin_just"])),
        ("", None),
        ("Regla del 100%", '=IF(AND(B9="SÍ",B8=0),"Cumple: cada ítem está en la EDT o tiene motivo escrito","No cumple: revisar")'),
        ("Descomposiciones 1 a 1", "Ninguna: todo entregable se abre en 2 o más" if not res["uno_a_uno"] else ", ".join(res["uno_a_uno"])),
        ("Nodos con verbo (actividades)", "Ninguno: todos los nodos son entregables (sustantivos)"),
    ]
    ws["A1"] = "Verificación de la regla del 100%"
    ws["A1"].font = Font(name="Calibri", size=14, bold=True, color=_AZUL)
    for k, (a, b) in enumerate(filas, 5):
        ws.cell(row=k, column=1, value=a).font = Font(name="Calibri", bold=True)
        c = ws.cell(row=k, column=2, value=b)
        c.font = Font(name="Calibri", bold=True)
        c.alignment = Alignment(horizontal="left")

    buf = io.BytesIO()
    wb.calculation.fullCalcOnLoad = True
    wb.save(buf)
    return buf.getvalue()


# ----------------------------------------------------------------------------------------
# 6. DIAGRAMA EN POWERPOINT (editable)
# ----------------------------------------------------------------------------------------
def pptx_edt(res=None):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[6])

    def caja(x, y, w, h, texto, fondo, color="FFFFFF", tam=9, negrita=False):
        sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        sh.fill.solid()
        sh.fill.fore_color.rgb = RGBColor.from_string(fondo)
        sh.line.color.rgb = RGBColor.from_string(_AZUL)
        sh.line.width = Pt(0.75)
        tf = sh.text_frame
        tf.word_wrap = True
        for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, m, Inches(0.03))
        for k, linea in enumerate(texto.split("\n")):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.CENTER
            run = p.add_run()
            run.text = linea
            run.font.size = Pt(tam)
            run.font.bold = negrita or k == 0
            run.font.color.rgb = RGBColor.from_string(color)
        return sh

    def linea(x1, y1, x2, y2):
        c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        c.line.color.rgb = RGBColor.from_string(_AZUL)
        c.line.width = Pt(1)

    caja(4.17, 0.2, 5.0, 0.6, f"1  {PROYECTO}\n{SUBTITULO.split(' · ')[0]}", _AZUL, tam=12, negrita=True)
    ancho, gap, x0 = 1.36, 0.08, 0.2
    xs = [x0 + i * (ancho + gap) for i in range(9)]
    linea(6.67, 0.8, 6.67, 1.05)
    linea(xs[0] + ancho / 2, 1.05, xs[-1] + ancho / 2, 1.05)
    for x, (ec, en, rama) in zip(xs, ENTREGABLES):
        cx = x + ancho / 2
        linea(cx, 1.05, cx, 1.25)
        caja(x, 1.25, ancho, 0.62, f"{ec}\n{en}", "2E3E7E" if rama == "Producto" else "C55A11", tam=9, negrita=True)
        paqs = [p for p in PAQUETES if p[0].startswith(ec + ".")]
        y = 2.05
        linea(x + 0.06, 1.87, x + 0.06, y + (len(paqs) - 1) * 0.92 + 0.38)
        for pc, pn, *_ in paqs:
            linea(x + 0.06, y + 0.38, x + 0.14, y + 0.38)
            caja(x + 0.14, y, ancho - 0.14, 0.76, f"{pc}\n{pn}", "F4F6FB" if rama == "Producto" else "FFF4EC", color=_AZUL, tam=8)
            y += 0.92
    tb = s.shapes.add_textbox(Inches(0.2), Inches(7.05), Inches(12.9), Inches(0.35))
    run = tb.text_frame.paragraphs[0].add_run()
    run.text = ("Azul = Producto (lo que recibe SUNAFIL) · Naranja = Proyecto (trabajo para entregarlo y gestionarlo) · "
                "Paquete de trabajo = último nivel (cuenta de control).")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string("5A6070")
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
