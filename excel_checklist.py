# -*- coding: utf-8 -*-
"""
==============================================================================
EXCEL DEL CHECKLIST  (Checklist_TDR_SUNAFIL.xlsx)
------------------------------------------------------------------------------
Hojas: 0 Resumen ejecutivo · A Ficha resumen · B Checklist puntuado
       B2 Factores de evaluación · C Mejoras · D Inconsistencias y riesgos
       E Recomendaciones · F Texto OCR crudo
Los conteos del resumen son fórmulas COUNTIF sobre la hoja B y el puntaje sale
de la hoja B2: si se corrige un veredicto o un puntaje a mano, el resumen se
actualiza solo.
==============================================================================
"""
import io, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

NOMBRE_ARCHIVO = "Checklist_TDR_SUNAFIL.xlsx"
NAVY, WHITE = "1F3864", "FFFFFF"
COLOR_VEREDICTO = {"SI": "C6EFCE", "NO": "FFC7CE", "PARCIAL": "FFEB9C",
                   "NO EVIDENCIADO": "D9D9D9", "NO APLICA": "F2F2F2"}
COLOR_PRIORIDAD = {"Alta": "FFC7CE", "Media": "FFEB9C", "Baja": "C6EFCE",
                   "Verificar en el PDF": "DDEBF7"}
MAX_CELDA = 32000   # Excel admite hasta 32 767 caracteres por celda

S0, SA, SB, SB2 = ("0. Resumen ejecutivo", "A. Ficha resumen", "B. Checklist puntuado",
                   "B2. Factores de evaluación")
SC, SD, SE, SF = "C. Mejoras", "D. Inconsistencias y riesgos", "E. Recomendaciones", "F. Texto OCR crudo"

COLS_B = ["N°", "Numeral TDR", "Título / Ítem", "Clase", "A) Alcance (Bases)", "Fuente (Bases)",
          "A+B) Consulta (Pliego)", "Valor exigido [A+B]", "A+B+C) Alcance final",
          "Valor ofertado", "Fuente (Propuesta)", "¿Cumple?", "Fuente del veredicto",
          "Motivo del veredicto", "Reflejo en el Contrato", "Mejora ofrecida", "Riesgo"]
ANCHOS_B = [6, 11, 32, 11, 40, 14, 34, 24, 40, 26, 16, 16, 12, 42, 32, 30, 30]
COLS_B2 = ["Factor", "Descripción", "Exigido", "Ofertado", "Fuente", "Puntaje máx.",
           "Puntaje obt.", "Motivo"]
ANCHOS_B2 = [12, 40, 36, 36, 18, 10, 10, 50]

_thin = Side(style="thin", color="BFBFBF")
BORDE = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
AJUSTE = Alignment(wrap_text=True, vertical="top")
CENTRO = Alignment(wrap_text=True, vertical="center", horizontal="center")


def _limpio(v):
    if isinstance(v, str):
        v = ILLEGAL_CHARACTERS_RE.sub("", v)[:MAX_CELDA]
    return v


def _celda(ws, r, c, v, formula=False):
    """Escribe un valor. Un texto que empieza con '=' se guarda como texto, no fórmula."""
    v = _limpio(v)
    cel = ws.cell(r, c, v)
    if isinstance(v, str) and v.startswith("=") and not formula:
        cel.data_type = "s"
    return cel


def _tabla(ws, encabezados, filas, anchos):
    """Encabezado azul oscuro, ajuste de texto, primera fila inmovilizada y filtros."""
    for j, h in enumerate(encabezados, 1):
        c = ws.cell(1, j, h)
        c.font = Font(bold=True, color=WHITE); c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = CENTRO; c.border = BORDE
    ws.row_dimensions[1].height = 32
    for i, fila in enumerate(filas, 2):
        for j, v in enumerate(fila, 1):
            c = _celda(ws, i, j, v); c.alignment = AJUSTE; c.border = BORDE
    for j, w in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"
    ultima = max(len(filas) + 1, 2)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(encabezados))}{ultima}"
    return ultima


def _paginas(textos):
    """Separa el texto de cada documento en (documento, página, texto)."""
    filas = []
    for doc, texto in textos.items():
        partes = re.split(r"^\[[^\]]+, p\. (\d+)\]\n", texto or "", flags=re.M)
        # partes = [antes, pag1, texto1, pag2, texto2, ...]
        for k in range(1, len(partes) - 1, 2):
            filas.append([doc, int(partes[k]), partes[k + 1].strip()])
    return filas


def construir_excel(resultado, textos):
    filas, factores, ficha = resultado["filas"], resultado["factores"], resultado["ficha"]
    mejoras, conteos = resultado["mejoras"], resultado["conteos"]
    wb = Workbook()

    # ---------------- B. Checklist puntuado ----------------
    wsB = wb.active; wsB.title = SB
    datos_b = [[f["n"], f["numeral"], f["titulo"], f.get("clase", ""), f.get("alcance_bases", ""),
                f.get("fuente_bases", ""), f.get("consulta_pliego", ""), f.get("valor_exigido", ""),
                f.get("alcance_final", ""), f.get("valor_ofertado", ""), f.get("fuente_propuesta", ""),
                f["veredicto"], f.get("fuente_veredicto", ""), f.get("motivo", ""),
                f.get("reflejo_contrato", ""), f.get("mejora", ""), f.get("riesgo", "")]
               for f in filas]
    ultima_b = _tabla(wsB, COLS_B, datos_b, ANCHOS_B)
    for r in range(2, ultima_b + 1):
        for c in (1, 2, 4, 12, 13):
            wsB.cell(r, c).alignment = CENTRO
        wsB.cell(r, 12).font = Font(bold=True)
    rango_l = f"L2:L{ultima_b}"
    dv = DataValidation(type="list", formula1='"SI,NO,PARCIAL,NO EVIDENCIADO,NO APLICA"',
                        allow_blank=True, showErrorMessage=True, errorTitle="Veredicto no válido",
                        error="Elige SI, NO, PARCIAL, NO EVIDENCIADO o NO APLICA.")
    wsB.add_data_validation(dv); dv.add(rango_l)
    for v, color in COLOR_VEREDICTO.items():   # el color sigue al valor si se edita a mano
        wsB.conditional_formatting.add(rango_l, CellIsRule(
            operator="equal", formula=[f'"{v}"'], fill=PatternFill("solid", fgColor=color)))

    # ---------------- B2. Factores de evaluación ----------------
    wsB2 = wb.create_sheet(SB2)
    datos_b2 = [[f["numeral"], f["titulo"] + (f" — {f['criterio']}" if f.get("criterio") else ""),
                 f.get("valor_exigido", ""), f.get("valor_ofertado", ""), f.get("fuente_propuesta", ""),
                 f["puntaje_max"], f["puntaje_obt"], f.get("motivo", "")] for f in factores]
    ultima_b2 = _tabla(wsB2, COLS_B2, datos_b2, ANCHOS_B2)
    for r in range(2, ultima_b2 + 1):
        for c in (1, 6, 7):
            wsB2.cell(r, c).alignment = CENTRO
        wsB2.cell(r, 7).font = Font(bold=True)
    fila_total = ultima_b2 + 1
    wsB2.cell(fila_total, 5, "TOTAL").font = Font(bold=True)
    for c, col in ((6, "F"), (7, "G")):
        t = wsB2.cell(fila_total, c, f"=SUM({col}2:{col}{ultima_b2})")
        t.font = Font(bold=True); t.alignment = CENTRO; t.border = BORDE

    # ---------------- 0. Resumen ejecutivo ----------------
    ws0 = wb.create_sheet(S0, 0)
    L = f"'{SB}'!L:L"
    fila_ficha = {x["campo"]: x["valor"] for x in ficha}
    riesgos = resultado["inconsistencias"]
    pmin = resultado["puntaje_minimo"]
    rango_pobt = f"'{SB2}'!G2:G{ultima_b2}"
    rango_pmax = f"'{SB2}'!F2:F{ultima_b2}"
    filas0 = [
        ("Veredicto global", f'=IF(COUNTIF({L},"NO")>0,"CON HALLAZGOS: "&COUNTIF({L},"NO")'
                             f'&" punto(s) en NO","SIN INCUMPLIMIENTOS DETECTADOS")', True),   # 2
        ("Entidad", fila_ficha.get("Entidad", ""), False),                                    # 3
        ("Objeto", fila_ficha.get("Objeto", ""), False),                                      # 4
        ("Contratista", fila_ficha.get("Contratista", ""), False),                            # 5
        ("Monto contractual", fila_ficha.get("Monto contractual", ""), False),                # 6
        ("Plazo de ejecución", fila_ficha.get("Plazo de ejecución", ""), False),              # 7
        ("Puntos del TDR rastreados (hoja B)", f"=COUNTA('{SB}'!A:A)-1", True),               # 8
        ("Acreditados con dato (SI)", f'=COUNTIF({L},"SI")', True),                           # 9
        ("Con hallazgo (NO)", f'=COUNTIF({L},"NO")', True),                                   # 10
        ("Parciales (PARCIAL)", f'=COUNTIF({L},"PARCIAL")', True),                            # 11
        ("Pendientes de evidencia (NO EVIDENCIADO)", f'=COUNTIF({L},"NO EVIDENCIADO")', True),  # 12
        ("No aplica según las Bases (NO APLICA)", f'=COUNTIF({L},"NO APLICA")', True),        # 13
        ("% de cumplimiento (SI / puntos que aplican)", "=IFERROR(B9/(B8-B13),0)", True),     # 14
        ("Cobertura documental (puntos que aplican con dato concreto)", conteos["cobertura"], False),  # 15
        ("Puntaje obtenido (hoja B2)", f"=SUM({rango_pobt})", True),                          # 16
        ("Puntaje máximo (hoja B2)", f"=SUM({rango_pmax})", True),                            # 17
        ("Puntaje mínimo (Bases)", pmin if pmin is not None else "No indicado en las Bases", False),  # 18
        ("¿Alcanza el puntaje mínimo?", '=IF(ISNUMBER(B18),IF(B16>=B18,"SÍ","NO"),"—")', True),  # 19
        ("Ítems en NO", "\n".join(conteos["items_no"]) or "Ninguno", False),
        ("Principales mejoras", "\n".join(f"{m['numerales']}: {m['mejora']}" for m in mejoras[:5])
                                or "Ninguna", False),
        ("Riesgos prioritarios", "\n".join(f"[{x['prioridad']}] {x['descripcion']}" for x in riesgos[:5])
                                 or "Ninguno", False),
        ("Filas 3.3.* quitadas por duplicar el TDR (3.4.*)",
         ", ".join(resultado.get("duplicados_quitados") or []) or "Ninguna", False),
    ]
    for j, h in enumerate(["Concepto", "Valor"], 1):
        c = ws0.cell(1, j, h)
        c.font = Font(bold=True, color=WHITE); c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = CENTRO; c.border = BORDE
    for i, (concepto, valor, es_formula) in enumerate(filas0, 2):
        a = ws0.cell(i, 1, concepto); a.font = Font(bold=True); a.alignment = AJUSTE; a.border = BORDE
        b = _celda(ws0, i, 2, valor, formula=es_formula); b.alignment = AJUSTE; b.border = BORDE
    ws0["B14"].number_format = "0.0%"   # % de cumplimiento
    ws0["B15"].number_format = "0.0%"   # cobertura
    ws0["B2"].font = Font(bold=True, size=12)
    ws0.column_dimensions["A"].width = 48; ws0.column_dimensions["B"].width = 90
    ws0.freeze_panes = "A2"

    # ---------------- A. Ficha resumen ----------------
    wsA = wb.create_sheet(SA, 1)
    _tabla(wsA, ["Campo", "Valor", "Fuente"],
           [[x["campo"], x["valor"], x["fuente"]] for x in ficha], [24, 70, 28])

    # ---------------- C. Mejoras (una fila por mejora) ----------------
    wsC = wb.create_sheet(SC)
    _tabla(wsC, ["Mejora", "Numeral", "Exigido", "Ofrecido", "Valor sobre el mínimo",
                 "Reflejo en el Contrato", "Evidencia"],
           [[m["mejora"], m["numerales"], m["exigido"], m["ofrecido"], m["valor_sobre_minimo"],
             m["reflejo_contrato"], m["evidencia"]] for m in mejoras],
           [40, 16, 26, 26, 20, 32, 18])

    # ---------------- D. Inconsistencias y riesgos ----------------
    wsD = wb.create_sheet(SD)
    _tabla(wsD, ["Prioridad", "Descripción", "Documentos en conflicto"],
           [[x["prioridad"], x["descripcion"], x["documentos"]] for x in riesgos], [13, 80, 40])
    for r, x in enumerate(riesgos, 2):
        c = wsD.cell(r, 1); c.alignment = CENTRO; c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor=COLOR_PRIORIDAD.get(x["prioridad"], "FFFFFF"))

    # ---------------- E. Recomendaciones ----------------
    wsE = wb.create_sheet(SE)
    _tabla(wsE, ["Ítem a monitorear", "Entregable que lo acredita", "Detalle"],
           [[x["item"], x["entregable"], x["detalle"]] for x in resultado["recomendaciones"]],
           [36, 36, 70])

    # ---------------- F. Texto OCR crudo ----------------
    wsF = wb.create_sheet(SF)
    _tabla(wsF, ["Documento", "Página", "Texto"], _paginas(textos), [12, 8, 120])

    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf.getvalue()
